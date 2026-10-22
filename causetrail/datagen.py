"""Synthetic, deliberately fragmented demo data for CauseTrail.

Produces four sources that look like what a small e-commerce ops team really has:

* ``orders.db``              SQLite  - orders, order_items, payments, refunds
* ``ad_spend.csv``           CSV     - daily spend per market / channel / campaign
* ``support_tickets.json``   JSON    - helpdesk export, references orders as "#<id>"
* ``inventory_supplier.xlsx`` Excel  - skus, warehouses, weekly supplier lead times

Planted root cause (only visible when you join across sources):
    Supplier SUP-07 (the sole Wearables supplier for the APAC warehouse SIN-1)
    slips its lead time from 7 to ~22 days from ISO week 36. APAC Wearables go on
    backorder from week 37, shipping delays balloon, customers open
    "where is my order" tickets, and a share of them ask for a refund. The refund
    table itself only says ``customer_request`` - the cause is upstream.

Red herrings, so the agent has something to falsify:
    * an APAC paid-social campaign ("APAC_Autumn_Push") ramps up in weeks 39-40;
    * a steady background of chargebacks and damaged-item refunds;
    * payment retries/failures at a flat rate.

Everything is driven by one ``random.Random(seed)`` so output is reproducible.
"""
from __future__ import annotations

import csv
import json
import random
import sqlite3
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

START = date(2026, 7, 13)  # Monday, ISO week 29
END = date(2026, 10, 4)  # Sunday, ISO week 40 (data "as of" this day)
REGIONS = {"NA": 40, "EU": 35, "APAC": 35, "LATAM": 15}  # mean orders/day
CHANNELS = {"organic": 0.45, "paid_search": 0.25, "paid_social": 0.18, "email": 0.12}
FAMILIES = {"Audio": 0.30, "Wearables": 0.30, "Home": 0.22, "Accessories": 0.18}
WAREHOUSES = {"NA": "DAL-1", "EU": "AMS-1", "APAC": "SIN-1", "LATAM": "GRU-1"}
SUPPLIERS = {
    "SUP-01": "Pacific Acoustics Ltd",
    "SUP-02": "Nordic Home Goods AB",
    "SUP-03": "Taipei Smart Devices Co",
    "SUP-04": "Guangzhou Cable Works",
    "SUP-07": "Shenzhen Lumen Electronics",
}
ROOT_SUPPLIER = "SUP-07"
ROOT_FAMILY = "Wearables"
ROOT_REGION = "APAC"

SKUS = [
    # sku, name, family, unit_price, unit_cost
    ("AUD-101", "Wireless Earbuds Lite", "Audio", 59.0, 22.0),
    ("AUD-102", "Over-ear ANC Headphones", "Audio", 149.0, 61.0),
    ("AUD-103", "Bluetooth Speaker Mini", "Audio", 39.0, 14.0),
    ("AUD-104", "Studio Monitor Pair", "Audio", 229.0, 98.0),
    ("WRB-201", "Fitness Band 3", "Wearables", 49.0, 17.0),
    ("WRB-202", "Smartwatch S", "Wearables", 179.0, 72.0),
    ("WRB-203", "Smartwatch S Pro", "Wearables", 249.0, 104.0),
    ("WRB-204", "Kids GPS Watch", "Wearables", 89.0, 33.0),
    ("HOM-301", "Smart Plug 2-pack", "Home", 29.0, 9.0),
    ("HOM-302", "LED Desk Lamp", "Home", 45.0, 16.0),
    ("HOM-303", "Air Quality Monitor", "Home", 119.0, 47.0),
    ("ACC-401", "USB-C Cable 2m", "Accessories", 15.0, 3.0),
    ("ACC-402", "65W GaN Charger", "Accessories", 42.0, 15.0),
    ("ACC-403", "Watch Strap Silicone", "Accessories", 19.0, 4.0),
]


def supplier_for(family: str, warehouse: str) -> str:
    if family == "Audio":
        return "SUP-01"
    if family == "Home":
        return "SUP-02"
    if family == "Accessories":
        return "SUP-04"
    return ROOT_SUPPLIER if warehouse == WAREHOUSES[ROOT_REGION] else "SUP-03"


def iso_week(d: date) -> int:
    return d.isocalendar()[1]


# Lead time of the root supplier by ISO week (promised is 7 days everywhere).
ROOT_LEAD = {36: 10, 37: 16, 38: 20, 39: 22, 40: 23}
# Probability that an APAC Wearables order goes on backorder, by ISO week.
BACKORDER_P = {37: 0.15, 38: 0.45, 39: 0.65, 40: 0.75}
CAMPAIGN_WEEKS = {39: 0.30, 40: 0.40}  # extra APAC paid-social demand

WISMO_SUBJECTS = [
    "Where is my order?",
    "Order still not shipped",
    "Any update on delivery?",
    "Order delayed - please advise",
    "Cancel if not shipped this week",
]
WISMO_BODIES = [
    "I ordered over a week ago and tracking still shows nothing.",
    "The site said 2-3 days. It has been much longer. What is going on?",
    "Please tell me when this will ship or refund me.",
    "Still waiting. I need this as a gift, otherwise I want my money back.",
]
OTHER_TICKETS = [
    ("How do I pair the device?", "setup", "Cannot get it to pair with my phone."),
    ("Size question", "presales", "Which strap size fits a 17cm wrist?"),
    ("Charged twice?", "billing", "I see two charges on my card statement."),
    ("Item arrived damaged", "damage", "The box was crushed and the item is cracked."),
    ("Change shipping address", "account", "Can I still change the address?"),
]


@dataclass
class Paths:
    root: Path

    @property
    def orders_db(self) -> Path:
        return self.root / "orders.db"

    @property
    def ad_spend(self) -> Path:
        return self.root / "ad_spend.csv"

    @property
    def tickets(self) -> Path:
        return self.root / "support_tickets.json"

    @property
    def inventory(self) -> Path:
        return self.root / "inventory_supplier.xlsx"


def _poisson(rng: random.Random, lam: float) -> int:
    # Knuth; lam is small enough here.
    import math

    L, k, p = math.exp(-lam), 0, 1.0
    while True:
        k += 1
        p *= rng.random()
        if p <= L:
            return k - 1


def _pick(rng: random.Random, weights: dict[str, float]) -> str:
    return rng.choices(list(weights), weights=list(weights.values()))[0]


def generate(out_dir: str | Path, seed: int = 7) -> Paths:
    rng = random.Random(seed)
    paths = Paths(Path(out_dir))
    paths.root.mkdir(parents=True, exist_ok=True)
    for p in (paths.orders_db, paths.ad_spend, paths.tickets, paths.inventory):
        if p.exists():
            p.unlink()

    skus_by_family: dict[str, list[tuple]] = {}
    for s in SKUS:
        skus_by_family.setdefault(s[2], []).append(s)

    orders, items, payments, refunds, tickets = [], [], [], [], []
    order_id, pay_id, refund_id, ticket_id = 100000, 1, 1, 50000
    as_of = END

    d = START
    while d <= END:
        wk = iso_week(d)
        weekend = d.weekday() >= 5
        for region, mean in REGIONS.items():
            lam = mean * (1.15 if weekend else 1.0)
            n = _poisson(rng, lam)
            extra_social = 0
            if region == ROOT_REGION and wk in CAMPAIGN_WEEKS:
                extra_social = _poisson(rng, mean * CAMPAIGN_WEEKS[wk])
            for i in range(n + extra_social):
                order_id += 1
                channel = "paid_social" if i >= n else _pick(rng, CHANNELS)
                family = _pick(rng, FAMILIES)
                sku = rng.choice(skus_by_family[family])
                qty = 1 if rng.random() < 0.85 else 2
                total = round(sku[3] * qty, 2)
                ordered_at = datetime.combine(d, datetime.min.time()) + timedelta(
                    minutes=rng.randint(0, 24 * 60 - 1)
                )

                delay = rng.randint(1, 3)
                backordered = (
                    region == ROOT_REGION
                    and family == ROOT_FAMILY
                    and rng.random() < BACKORDER_P.get(wk, 0.01)
                )
                if backordered:
                    delay += rng.randint(8, 18)
                ship_day = d + timedelta(days=delay)
                shipped = ship_day if ship_day <= as_of else None

                # refunds -------------------------------------------------
                refund = None
                if backordered and rng.random() < 0.45:
                    r_day = d + timedelta(days=rng.randint(7, 13))
                    reason = "customer_request" if rng.random() < 0.85 else "other"
                    refund = (r_day, reason)
                elif rng.random() < 0.03:
                    reason = rng.choices(
                        ["customer_request", "damaged", "not_as_described", "chargeback", "other"],
                        weights=[0.45, 0.2, 0.15, 0.12, 0.08],
                    )[0]
                    r_day = d + timedelta(days=delay + rng.randint(3, 12))
                    refund = (r_day, reason)
                if refund and refund[0] > as_of:
                    refund = None

                if refund:
                    status = "refunded"
                elif shipped is None:
                    status = "backordered" if backordered else "processing"
                else:
                    status = "shipped"

                orders.append(
                    (
                        order_id,
                        ordered_at.isoformat(sep=" "),
                        f"C{rng.randint(10000, 99999)}",
                        region,
                        channel,
                        (d + timedelta(days=3)).isoformat(),
                        shipped.isoformat() if shipped else None,
                        status,
                    )
                )
                items.append((order_id, sku[0], qty, sku[3], sku[4]))

                # payments: flat ~3% first-attempt failure, then a capture
                method = rng.choices(["card", "paypal", "bnpl"], weights=[0.6, 0.27, 0.13])[0]
                if rng.random() < 0.03:
                    payments.append((pay_id, order_id, ordered_at.isoformat(sep=" "), total, method, "failed"))
                    pay_id += 1
                payments.append(
                    (pay_id, order_id, (ordered_at + timedelta(minutes=2)).isoformat(sep=" "), total, method, "captured")
                )
                pay_id += 1

                if refund:
                    refunds.append((refund_id, order_id, refund[0].isoformat(), total, refund[1]))
                    refund_id += 1

                # support tickets ------------------------------------------
                if backordered and rng.random() < 0.75:
                    t_day = d + timedelta(days=rng.randint(4, 8))
                    if t_day <= as_of:
                        ticket_id += 1
                        tickets.append(
                            {
                                "ticket_id": f"T-{ticket_id}",
                                "created_at": f"{t_day.isoformat()}T{rng.randint(8, 21):02d}:{rng.randint(0, 59):02d}:00Z",
                                "order_ref": f"#{order_id}",
                                "subject": rng.choice(WISMO_SUBJECTS),
                                "body": rng.choice(WISMO_BODIES),
                                "tags": rng.choice([["shipping"], ["shipping", "urgent"], [], ["delivery"]]),
                                "priority": rng.choice(["normal", "normal", "high"]),
                                "requester": {"customer_id": orders[-1][2], "lang": "en"},
                            }
                        )
                if rng.random() < 0.05 or (refund and refund[1] == "damaged" and rng.random() < 0.6):
                    subj, tag, body = (
                        OTHER_TICKETS[3] if refund and refund[1] == "damaged" else rng.choice(OTHER_TICKETS)
                    )
                    t_day = d + timedelta(days=rng.randint(0, delay + 4))
                    if t_day <= as_of:
                        ticket_id += 1
                        tickets.append(
                            {
                                "ticket_id": f"T-{ticket_id}",
                                "created_at": f"{t_day.isoformat()}T{rng.randint(8, 21):02d}:{rng.randint(0, 59):02d}:00Z",
                                "order_ref": f"#{order_id}",
                                "subject": subj,
                                "body": body,
                                "tags": [tag],
                                "priority": "normal",
                                "requester": {"customer_id": orders[-1][2], "lang": "en"},
                            }
                        )
        d += timedelta(days=1)

    # ---- orders.db -----------------------------------------------------------
    con = sqlite3.connect(paths.orders_db)
    con.executescript(
        """
        CREATE TABLE orders (order_id INTEGER PRIMARY KEY, order_ts TEXT, customer_id TEXT,
            region TEXT, channel TEXT, promised_ship_date TEXT, shipped_date TEXT, status TEXT);
        CREATE TABLE order_items (order_id INTEGER, sku TEXT, qty INTEGER, unit_price REAL, unit_cost REAL);
        CREATE TABLE payments (payment_id INTEGER PRIMARY KEY, order_id INTEGER, attempted_at TEXT,
            amount REAL, method TEXT, status TEXT);
        CREATE TABLE refunds (refund_id INTEGER PRIMARY KEY, order_id INTEGER, refunded_at TEXT,
            amount REAL, reason_code TEXT);
        """
    )
    con.executemany("INSERT INTO orders VALUES (?,?,?,?,?,?,?,?)", orders)
    con.executemany("INSERT INTO order_items VALUES (?,?,?,?,?)", items)
    con.executemany("INSERT INTO payments VALUES (?,?,?,?,?,?)", payments)
    con.executemany("INSERT INTO refunds VALUES (?,?,?,?,?)", refunds)
    con.commit()
    con.close()

    # ---- ad_spend.csv --------------------------------------------------------
    with paths.ad_spend.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["date", "market", "channel", "campaign", "spend_usd", "impressions", "clicks"])
        d = START
        while d <= END:
            wk = iso_week(d)
            for region in REGIONS:
                for ch, base in (("paid_search", 420.0), ("paid_social", 310.0)):
                    scale = REGIONS[region] / 25
                    spend = base * scale * rng.uniform(0.85, 1.15)
                    campaign = f"{region}_Always_On_{ch}"
                    rows = [(campaign, spend)]
                    if region == ROOT_REGION and ch == "paid_social" and wk in CAMPAIGN_WEEKS:
                        rows.append(("APAC_Autumn_Push", 2600.0 * CAMPAIGN_WEEKS[wk] / 0.3 * rng.uniform(0.9, 1.1)))
                    for camp, sp in rows:
                        impressions = int(sp * rng.uniform(90, 130))
                        w.writerow([d.isoformat(), region, ch, camp, round(sp, 2), impressions,
                                    int(impressions * rng.uniform(0.008, 0.02))])
            d += timedelta(days=1)

    # ---- support_tickets.json ------------------------------------------------
    tickets.sort(key=lambda t: t["created_at"])
    paths.tickets.write_text(json.dumps({"exported_at": f"{END.isoformat()}T23:59:00Z", "tickets": tickets}, indent=1))

    # ---- inventory_supplier.xlsx --------------------------------------------
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "skus"
    ws.append(["sku", "product_name", "family", "list_price", "unit_cost"])
    for s in SKUS:
        ws.append(list(s))
    ws2 = wb.create_sheet("warehouses")
    ws2.append(["warehouse", "serves_region", "city"])
    for region, wh, city in (("NA", "DAL-1", "Dallas"), ("EU", "AMS-1", "Amsterdam"),
                             ("APAC", "SIN-1", "Singapore"), ("LATAM", "GRU-1", "Sao Paulo")):
        ws2.append([wh, region, city])
    ws3 = wb.create_sheet("suppliers")
    ws3.append(["supplier_id", "supplier_name"])
    for k, v in SUPPLIERS.items():
        ws3.append([k, v])
    ws4 = wb.create_sheet("lead_times")
    ws4.append(["week_start", "warehouse", "family", "supplier_id", "promised_lead_days",
                "actual_lead_days", "on_hand_units", "inbound_units"])
    wk_start = START
    while wk_start <= END:
        wk = iso_week(wk_start)
        for region, wh in WAREHOUSES.items():
            for fam in FAMILIES:
                sup = supplier_for(fam, wh)
                actual = 7 + rng.choice([-1, 0, 0, 0, 1, 1, 2])
                on_hand = int(REGIONS[region] * FAMILIES[fam] * 7 * rng.uniform(1.6, 2.4))
                inbound = int(on_hand * rng.uniform(0.8, 1.2))
                if sup == ROOT_SUPPLIER and wk in ROOT_LEAD:
                    actual = ROOT_LEAD[wk] + rng.choice([-1, 0, 1])
                    depletion = {36: 0.8, 37: 0.35, 38: 0.12, 39: 0.05, 40: 0.03}[wk]
                    on_hand = int(on_hand * depletion)
                    inbound = int(inbound * depletion)
                ws4.append([wk_start, wh, fam, sup, 7, actual, on_hand, inbound])
        wk_start += timedelta(days=7)
    wb.save(paths.inventory)
    return paths


if __name__ == "__main__":  # pragma: no cover
    import sys

    p = generate(sys.argv[1] if len(sys.argv) > 1 else "data")
    print(f"wrote demo data to {p.root}")
